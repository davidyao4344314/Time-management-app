"""Parser/storage tests use synthetic documents and disposable SQLite only."""

from io import BytesIO
import sqlite3
import unittest

from backend.app.files.contracts import CHUNK_CHARS, MAX_UPLOAD_BYTES, FileImportError
from backend.app.files.ingestion import ingest_file
from backend.app.files.parser import parse_file
from backend.app.files import storage


class FileIngestionTests(unittest.TestCase):
    def setUp(self):
        self.connection = sqlite3.connect(":memory:")
        self.addCleanup(self.connection.close)
        self.connection.execute("PRAGMA foreign_keys=ON")
        storage.migrate(self.connection)

    def test_text_roundtrip_order_and_duplicate_filename_identity(self):
        raw = ("Recursion assignment\r\n" + "Question 4: use a base case.\n" * 300).encode()
        first = ingest_file(self.connection, "owner", "assignment2.txt", raw)
        second = ingest_file(self.connection, "owner", "assignment2.txt", raw)
        self.assertNotEqual(first["file_id"], second["file_id"])
        stored = storage.get_file(self.connection, "owner", first["file_id"])
        self.assertEqual("".join(chunk["text"] for chunk in stored["chunks"]), stored["text"])
        self.assertTrue(all(len(chunk["text"]) <= CHUNK_CHARS for chunk in stored["chunks"]))
        self.assertEqual([c["position"] for c in stored["chunks"]], list(range(1, len(stored["chunks"]) + 1)))
        self.assertNotIn("text", first)
        self.assertIsNone(storage.get_file(self.connection, "another-owner", first["file_id"]))
        self.assertEqual(storage.list_files(self.connection, "another-owner"), [])
        storage.migrate(self.connection)
        self.assertEqual(len(storage.list_files(self.connection, "owner")), 2)

    def test_docx_extracts_paragraphs_and_table_without_fake_page_numbers(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Recursion instructions")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Question 4"
        output = BytesIO()
        document.save(output)
        result = parse_file("assignment.docx", output.getvalue())
        self.assertIn("Question 4", result["text"])
        self.assertTrue(all(chunk["page"] is None for chunk in result["chunks"]))

    def test_pdf_extracts_text_and_real_page_numbers(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 30 700 Td (Question 4: recursion.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        data = BytesIO()
        writer.write(data)
        result = parse_file("assignment2.pdf", data.getvalue())
        self.assertIn("recursion", result["text"])
        self.assertEqual(result["chunks"][0]["page"], 1)

    def test_empty_unsupported_invalid_paths_and_size_fail_before_saving(self):
        cases = [("empty.txt", b""), ("archive.zip", b"data"), ("../secret.txt", b"text"),
                 ("/etc/file.txt", b"text"), ("bad.pdf", b"not a PDF"),
                 ("bad.docx", b"not a ZIP"), ("bad.txt", b"\xff"),
                 ("huge.txt", b"a" * (MAX_UPLOAD_BYTES + 1)),
                 ("huge-text.txt", b"a" * 250_001)]
        for filename, data in cases:
            with self.subTest(filename=filename), self.assertRaises(FileImportError):
                ingest_file(self.connection, "owner", filename, data)
        self.assertEqual(storage.list_files(self.connection, "owner"), [])


if __name__ == "__main__":
    unittest.main()
